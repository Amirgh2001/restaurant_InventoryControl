"""
Reports, Analytics, Dashboards, and Excel Exporter
Provides management insights, discrepancy reports, user performance,
abnormal consumption/waste detectors, and native Excel (.xlsx) file generation.
"""

import io
import csv
from datetime import datetime, timedelta
from typing import Dict, Any, List, Optional
from .database import get_connection
from .inventory_service import get_stock_matrix, get_low_stock_alerts

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

def gregorian_to_jalali(gy: int, gm: int, gd: int):
    """Converts a Gregorian date (year, month, day) to Solar Hijri (Jalali)."""
    g_d_m = [0, 31, 59, 90, 120, 151, 181, 212, 243, 273, 304, 334]
    if gy > 1600:
        jy = 979
        gy -= 1600
    else:
        jy = 0
        gy -= 621
    gy2 = gy + 1 if gm > 2 else gy
    days = 365 * gy + (gy2 + 3) // 4 - (gy2 + 99) // 100 + (gy2 + 399) // 400 - 80 + gd + g_d_m[gm - 1]
    jy += 33 * (days // 12053)
    days %= 12053
    jy += 4 * (days // 1461)
    days %= 1461
    if days > 365:
        jy += (days - 1) // 365
        days = (days - 1) % 365
    if days < 186:
        jm = 1 + days // 31
        jd = 1 + days % 31
    else:
        jm = 7 + (days - 186) // 30
        jd = 1 + (days - 186) % 30
    return jy, jm, jd

def format_datetime_persian(dt_str: Optional[str], include_time: bool = True) -> str:
    """Converts ISO or standard datetime string to Persian Solar Hijri formatted string."""
    if not dt_str:
        return "-"
    try:
        clean = dt_str.replace("T", " ")
        parts = clean.split(" ")
        date_parts = [int(p) for p in parts[0].split("-")]
        jy, jm, jd = gregorian_to_jalali(date_parts[0], date_parts[1], date_parts[2])
        persian_date = f"{jy:04d}/{jm:02d}/{jd:02d}"
        if include_time and len(parts) > 1:
            time_part = parts[1][:5]
            return f"{persian_date} - {time_part}"
        return persian_date
    except Exception:
        return dt_str

def get_dashboard_summary() -> Dict[str, Any]:
    """Generates high-level executive KPI metrics for restaurant managers."""
    conn = get_connection()
    cursor = conn.cursor()

    # 1. Total items count
    cursor.execute("SELECT COUNT(*) FROM items WHERE is_active = 1;")
    total_items = cursor.fetchone()[0]

    # 2. Total locations count
    cursor.execute("SELECT COUNT(*) FROM locations WHERE is_active = 1;")
    total_locations = cursor.fetchone()[0]

    # 3. Low stock count
    low_stock_items = get_low_stock_alerts()
    low_stock_count = len(low_stock_items)

    # 4. Total transactions today
    today_str = datetime.now().strftime("%Y-%m-%d")
    cursor.execute("""
        SELECT COUNT(*) FROM inventory_transactions 
        WHERE created_at LIKE ?;
    """, (f"{today_str}%",))
    tx_today_count = cursor.fetchone()[0]

    # 5. Waste in last 7 days grouped by reasons
    week_ago = (datetime.now() - timedelta(days=7)).isoformat()
    cursor.execute("""
        SELECT wr.title AS reason, COUNT(it.id) AS count, SUM(it.quantity_base) AS total_qty
        FROM inventory_transactions it
        JOIN waste_reasons wr ON it.reason_id = wr.id
        WHERE it.transaction_type = 'WASTE' AND it.created_at >= ?
        GROUP BY wr.title
        ORDER BY total_qty DESC;
    """, (week_ago,))
    waste_breakdown = [dict(r) for r in cursor.fetchall()]

    # 6. Top consumed items in last 7 days
    cursor.execute("""
        SELECT i.name AS item_name, u.name AS unit_name, SUM(it.quantity_base) AS total_consumed
        FROM inventory_transactions it
        JOIN items i ON it.item_id = i.id
        JOIN units u ON i.base_unit_id = u.id
        WHERE it.transaction_type = 'CONSUMPTION' AND it.created_at >= ?
        GROUP BY it.item_id
        ORDER BY total_consumed DESC
        LIMIT 5;
    """, (week_ago,))
    top_consumed = [dict(r) for r in cursor.fetchall()]

    # 7. Recent 10 transactions
    cursor.execute("""
        SELECT 
            it.id, it.transaction_type, it.quantity, it.quantity_base, it.created_at, it.notes,
            i.name AS item_name, u.symbol AS unit_symbol,
            fl.name AS from_location_name, tl.name AS to_location_name,
            u_user.full_name AS user_name, wr.title AS waste_reason
        FROM inventory_transactions it
        JOIN items i ON it.item_id = i.id
        JOIN units u ON it.unit_id = u.id
        JOIN users u_user ON it.user_id = u_user.id
        LEFT JOIN locations fl ON it.from_location_id = fl.id
        LEFT JOIN locations tl ON it.to_location_id = tl.id
        LEFT JOIN waste_reasons wr ON it.reason_id = wr.id
        ORDER BY it.id DESC
        LIMIT 10;
    """)
    recent_transactions = [dict(r) for r in cursor.fetchall()]

    # 8. Pending stocktake sessions
    cursor.execute("""
        SELECT ss.id, ss.title, ss.status, ss.created_at, l.name AS location_name, u.full_name AS creator_name
        FROM stocktake_sessions ss
        JOIN locations l ON ss.location_id = l.id
        JOIN users u ON ss.created_by = u.id
        WHERE ss.status IN ('DRAFT', 'SUBMITTED')
        ORDER BY ss.id DESC;
    """)
    pending_stocktakes = [dict(r) for r in cursor.fetchall()]

    conn.close()

    return {
        "total_items": total_items,
        "total_locations": total_locations,
        "low_stock_count": low_stock_count,
        "tx_today_count": tx_today_count,
        "waste_breakdown": waste_breakdown,
        "top_consumed": top_consumed,
        "recent_transactions": recent_transactions,
        "pending_stocktakes": pending_stocktakes,
        "low_stock_alerts": low_stock_items
    }

def get_filtered_transactions(filters: Dict[str, Any]) -> List[Dict[str, Any]]:
    """Fetches transactions with extensive filtering capabilities."""
    conn = get_connection()
    cursor = conn.cursor()

    query = """
        SELECT 
            it.id, it.transaction_type, it.quantity, it.quantity_base, it.created_at, it.notes,
            it.batch_number, it.expiry_date,
            i.id AS item_id, i.name AS item_name, c.name AS category_name,
            u.name AS unit_name, u.symbol AS unit_symbol,
            fl.id AS from_location_id, fl.name AS from_location_name,
            tl.id AS to_location_id, tl.name AS to_location_name,
            u_user.id AS user_id, u_user.full_name AS user_name,
            wr.title AS waste_reason
        FROM inventory_transactions it
        JOIN items i ON it.item_id = i.id
        JOIN categories c ON i.category_id = c.id
        JOIN units u ON it.unit_id = u.id
        JOIN users u_user ON it.user_id = u_user.id
        LEFT JOIN locations fl ON it.from_location_id = fl.id
        LEFT JOIN locations tl ON it.to_location_id = tl.id
        LEFT JOIN waste_reasons wr ON it.reason_id = wr.id
        WHERE 1=1
    """
    params = []

    if filters.get("transaction_type"):
        query += " AND it.transaction_type = ?"
        params.append(filters["transaction_type"])

    if filters.get("item_id"):
        query += " AND it.item_id = ?"
        params.append(int(filters["item_id"]))

    if filters.get("category_id"):
        query += " AND i.category_id = ?"
        params.append(int(filters["category_id"]))

    if filters.get("location_id"):
        loc_id = int(filters["location_id"])
        query += " AND (it.from_location_id = ? OR it.to_location_id = ?)"
        params.extend([loc_id, loc_id])

    if filters.get("user_id"):
        query += " AND it.user_id = ?"
        params.append(int(filters["user_id"]))

    if filters.get("start_date"):
        query += " AND it.created_at >= ?"
        params.append(filters["start_date"])

    if filters.get("end_date"):
        query += " AND it.created_at <= ?"
        params.append(filters["end_date"])

    query += " ORDER BY it.id DESC LIMIT 500;"

    cursor.execute(query, params)
    results = [dict(r) for r in cursor.fetchall()]
    conn.close()
    return results

def get_stocktake_details(session_id: int) -> Dict[str, Any]:
    """Retrieves full details of a stocktake session including discrepancy items."""
    conn = get_connection()
    cursor = conn.cursor()

    cursor.execute("""
        SELECT 
            ss.id, ss.title, ss.status, ss.notes, ss.created_at, ss.approved_at,
            l.name AS location_name, l.id AS location_id,
            uc.full_name AS creator_name,
            ua.full_name AS approver_name
        FROM stocktake_sessions ss
        JOIN locations l ON ss.location_id = l.id
        JOIN users uc ON ss.created_by = uc.id
        LEFT JOIN users ua ON ss.approved_by = ua.id
        WHERE ss.id = ?;
    """, (session_id,))
    session = cursor.fetchone()
    if not session:
        conn.close()
        return {}

    cursor.execute("""
        SELECT 
            si.id, si.item_id, si.system_quantity, si.counted_quantity, si.difference, si.notes,
            i.name AS item_name, i.code AS item_code, c.name AS category_name, u.name AS unit_name
        FROM stocktake_items si
        JOIN items i ON si.item_id = i.id
        JOIN categories c ON i.category_id = c.id
        JOIN units u ON i.base_unit_id = u.id
        WHERE si.session_id = ?
        ORDER BY ABS(si.difference) DESC, i.name;
    """, (session_id,))
    items = [dict(r) for r in cursor.fetchall()]

    conn.close()
    return {
        "session": dict(session),
        "items": items
    }

def get_user_performance_report() -> List[Dict[str, Any]]:
    """Analyzes operator and staff log activity breakdown."""
    conn = get_connection()
    cursor = conn.cursor()

    cursor.execute("""
        SELECT 
            u.id, u.full_name, u.username, u.role,
            COUNT(it.id) AS total_transactions,
            SUM(CASE WHEN it.transaction_type = 'INBOUND' THEN 1 ELSE 0 END) AS inbound_count,
            SUM(CASE WHEN it.transaction_type = 'TRANSFER' THEN 1 ELSE 0 END) AS transfer_count,
            SUM(CASE WHEN it.transaction_type = 'CONSUMPTION' THEN 1 ELSE 0 END) AS consumption_count,
            SUM(CASE WHEN it.transaction_type = 'WASTE' THEN 1 ELSE 0 END) AS waste_count,
            MAX(it.created_at) AS last_activity
        FROM users u
        LEFT JOIN inventory_transactions it ON u.id = it.user_id
        WHERE u.is_active = 1
        GROUP BY u.id
        ORDER BY total_transactions DESC;
    """)
    perf = [dict(r) for r in cursor.fetchall()]
    conn.close()
    return perf

def get_abnormal_waste_and_consumption() -> Dict[str, Any]:
    """Identifies items with unusually high waste ratios or unexpected spikes."""
    conn = get_connection()
    cursor = conn.cursor()

    # Compare waste vs consumption over last 14 days
    two_weeks_ago = (datetime.now() - timedelta(days=14)).isoformat()

    cursor.execute("""
        SELECT 
            i.id, i.name AS item_name, c.name AS category_name, u.name AS unit_name,
            COALESCE(SUM(CASE WHEN it.transaction_type = 'CONSUMPTION' THEN it.quantity_base ELSE 0 END), 0) AS total_consumed,
            COALESCE(SUM(CASE WHEN it.transaction_type = 'WASTE' THEN it.quantity_base ELSE 0 END), 0) AS total_waste
        FROM items i
        JOIN categories c ON i.category_id = c.id
        JOIN units u ON i.base_unit_id = u.id
        LEFT JOIN inventory_transactions it ON i.id = it.item_id AND it.created_at >= ?
        GROUP BY i.id
        HAVING total_waste > 0 OR total_consumed > 0
        ORDER BY total_waste DESC;
    """, (two_weeks_ago,))
    rows = cursor.fetchall()

    abnormal_items = []
    for r in rows:
        consumed = r["total_consumed"]
        waste = r["total_waste"]
        total = consumed + waste
        waste_percent = (waste / total * 100.0) if total > 0 else 0.0

        # Flag items where waste exceeds 5% of usage or waste is unusually high
        is_abnormal = waste_percent > 5.0 or waste > 5.0

        abnormal_items.append({
            "item_name": r["item_name"],
            "category_name": r["category_name"],
            "unit_name": r["unit_name"],
            "total_consumed": round(consumed, 2),
            "total_waste": round(waste, 2),
            "waste_percentage": round(waste_percent, 1),
            "is_abnormal": is_abnormal
        })

    conn.close()
    return {"analysis": abnormal_items}

# ----------------- Excel (.xlsx) Generation -----------------

def export_stock_to_excel() -> io.BytesIO:
    """Generates a beautifully styled RTL Excel spreadsheet of the current inventory."""
    output = io.BytesIO()

    if OPENPYXL_AVAILABLE:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "موجودی لحظه‌ای انبارها"
        ws.views.sheetView[0].rightToLeft = True  # Enable Persian RTL layout in Excel!

        # Styles
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Navy Blue
        header_font = Font(name="Tahoma", size=11, bold=True, color="FFFFFF")
        alert_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid") # Soft Red
        alert_font = Font(name="Tahoma", size=10, bold=True, color="991B1B")
        normal_font = Font(name="Tahoma", size=10)
        center_align = Alignment(horizontal="center", vertical="center")
        right_align = Alignment(horizontal="right", vertical="center")
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        matrix = get_stock_matrix()

        # Get location names for dynamic columns
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT id, name FROM locations WHERE is_active = 1 ORDER BY id;")
        locations = cursor.fetchall()
        conn.close()

        # Headers
        headers = ["ردیف", "کد کالا", "نام کالا", "دسته‌بندی", "واحد"]
        for loc in locations:
            headers.append(f"موجودی {loc['name']}")
        headers.extend(["مجموع موجودی کل", "حداقل مجاز هشدار", "وضعیت"])

        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border
        ws.row_dimensions[1].height = 28

        # Data rows
        for idx, item in enumerate(matrix, 1):
            status = "بحرانی (کمبود)" if (item["is_overall_low"] or item["has_location_alert"]) else "نرمال"
            row_data = [
                idx,
                item["code"] or "-",
                item["name"],
                item["category_name"],
                item["unit_name"]
            ]
            for loc in locations:
                loc_stock = item["location_stocks"].get(loc["id"], {}).get("quantity", 0.0)
                row_data.append(loc_stock)

            row_data.extend([
                item["total_quantity"],
                item["min_stock_alert"],
                status
            ])

            ws.append(row_data)
            row_num = idx + 1
            ws.row_dimensions[row_num].height = 22

            is_alert = item["is_overall_low"] or item["has_location_alert"]
            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=row_num, column=col_idx)
                cell.font = alert_font if is_alert else normal_font
                cell.border = thin_border
                if col_idx in (1, 2, 4, 5, len(row_data)):
                    cell.alignment = center_align
                else:
                    cell.alignment = right_align

                if is_alert and col_idx == len(row_data):
                    cell.fill = alert_fill

        # Auto column width
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        wb.save(output)
        output.seek(0)
    else:
        # Fallback to UTF-8 BOM CSV
        output.write(b'\xef\xbb\xbf') # UTF-8 BOM for Persian Excel compatibility
        writer = csv.writer(io.TextIOWrapper(output, encoding='utf-8', write_through=True))
        matrix = get_stock_matrix()
        writer.writerow(["کد کالا", "نام کالا", "دسته‌بندی", "مجموع کل", "واحد", "حداقل هشدار"])
        for itm in matrix:
            writer.writerow([itm["code"], itm["name"], itm["category_name"], itm["total_quantity"], itm["unit_name"], itm["min_stock_alert"]])
        output.seek(0)

    return output

def export_transactions_to_excel(filters: Dict[str, Any]) -> io.BytesIO:
    """Exports filtered transaction log to formatted Excel file."""
    output = io.BytesIO()
    txs = get_filtered_transactions(filters)

    type_labels = {
        "INBOUND": "ورود کالا",
        "TRANSFER": "انتقال بین‌بخشی",
        "CONSUMPTION": "ثبت مصرف",
        "WASTE": "ثبت ضایعات",
        "ADJUSTMENT_ADD": "اصلاح انبارگردانی (افزایش)",
        "ADJUSTMENT_SUB": "اصلاح انبارگردانی (کاهش)"
    }

    if OPENPYXL_AVAILABLE:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "گزارش تراکنش‌های انبار"
        ws.views.sheetView[0].rightToLeft = True

        header_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid") # Dark Teal
        header_font = Font(name="Tahoma", size=11, bold=True, color="FFFFFF")
        normal_font = Font(name="Tahoma", size=10)
        center_align = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        headers = ["شناسه", "نوع عملیات", "کالا", "دسته‌بندی", "مقدار", "واحد", "مبدأ", "مقصد", "دلیل ضایعات", "کاربر ثبت‌کننده", "تاریخ و ساعت", "توضیحات"]
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border
        ws.row_dimensions[1].height = 26

        for idx, tx in enumerate(txs, 1):
            row_data = [
                tx["id"],
                type_labels.get(tx["transaction_type"], tx["transaction_type"]),
                tx["item_name"],
                tx["category_name"],
                tx["quantity"],
                tx["unit_symbol"],
                tx["from_location_name"] or "-",
                tx["to_location_name"] or "-",
                tx["waste_reason"] or "-",
                tx["user_name"],
                format_datetime_persian(tx["created_at"], include_time=True),
                tx["notes"] or "-"
            ]
            ws.append(row_data)
            row_num = idx + 1
            ws.row_dimensions[row_num].height = 20
            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=row_num, column=col_idx)
                cell.font = normal_font
                cell.border = thin_border
                cell.alignment = center_align

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        wb.save(output)
        output.seek(0)
    else:
        output.write(b'\xef\xbb\xbf')
        writer = csv.writer(io.TextIOWrapper(output, encoding='utf-8', write_through=True))
        writer.writerow(["شناسه", "نوع عملیات", "کالا", "مقدار", "مبدأ", "مقصد", "کاربر", "تاریخ شمسی"])
        for tx in txs:
            writer.writerow([tx["id"], type_labels.get(tx["transaction_type"], tx["transaction_type"]), tx["item_name"], tx["quantity"], tx["from_location_name"], tx["to_location_name"], tx["user_name"], format_datetime_persian(tx["created_at"], include_time=True)])
        output.seek(0)

    return output
